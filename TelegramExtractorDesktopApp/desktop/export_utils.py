# desktop/export_utils.py
from __future__ import annotations
import pandas as pd
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas
from reportlab.lib.units import mm
from docx import Document
from docx.shared import Inches

def export_csv(df: pd.DataFrame, path: str):
    df.to_csv(path, index=False, encoding="utf-8-sig")

def export_excel(df: pd.DataFrame, path: str):
    with pd.ExcelWriter(path, engine="xlsxwriter") as w:
        df.to_excel(w, index=False, sheet_name="Data")
        ws = w.sheets["Data"]
        for i, col in enumerate(df.columns):
            width = max(12, min(60, int(df[col].astype(str).str.len().mean()) + 6))
            ws.set_column(i, i, width)

def export_pdf(df: pd.DataFrame, path: str, title: str = "Export"):
    c = canvas.Canvas(path, pagesize=landscape(A4))
    width, height = landscape(A4)
    margin = 12 * mm
    y = height - margin
    c.setFont("Helvetica-Bold", 14)
    c.drawString(margin, y, title)
    y -= 10 * mm
    c.setFont("Helvetica", 9)

    cols = list(df.columns)
    col_w = (width - 2 * margin) / max(1, len(cols))
    # header
    for i, col in enumerate(cols):
        c.drawString(margin + i * col_w + 2, y, str(col)[:32])
    y -= 6 * mm
    c.line(margin, y + 2, width - margin, y + 2)
    y -= 4 * mm

    # rows
    for _, row in df.iterrows():
        if y < margin:
            c.showPage(); y = height - margin
            c.setFont("Helvetica", 9)
        for i, col in enumerate(cols):
            txt = str(row.get(col, ""))[:64]
            c.drawString(margin + i * col_w + 2, y, txt)
        y -= 6 * mm
    c.save()

def export_docx(df: pd.DataFrame, path: str, title: str = "Export"):
    doc = Document()
    doc.add_heading(title, level=1)
    table = doc.add_table(rows=1, cols=len(df.columns))
    hdr = table.rows[0].cells
    for i, col in enumerate(df.columns):
        hdr[i].text = str(col)
    for _, r in df.iterrows():
        row_cells = table.add_row().cells
        for i, col in enumerate(df.columns):
            row_cells[i].text = str(r.get(col, ""))
    doc.save(path)
